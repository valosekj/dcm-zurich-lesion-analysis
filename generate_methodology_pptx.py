"""Generate the editable methodology PPTX for format_sct_analyze_lesion_table_max-per-slice.py.

The deck is a 5-slide progressive build (each slide reveals one more step of the pipeline):

    slide 0: title + per-slice input table (no highlights) + legend + C4/C5 level labels
    slide 1: + per-level-max highlights in the input table + step 1
    slide 2: + per-level maxima table (blue/orange) + arrow
    slide 3: + green "selected max across levels" winners + step 2 + arrow
    slide 4: + per-side output box + arrow

Run with the dedicated `pptx` conda env (python-pptx not installed elsewhere):
    conda run -n pptx python generate_methodology_pptx.py
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

# ---- palette ----
HEADER = RGBColor(0x2B, 0x3A, 0x67)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1B, 0x24, 0x30)
MUTE = RGBColor(0x5B, 0x66, 0x73)
LEFT = RGBColor(0xCF, 0xE3, 0xFF)   # max left within level (blue)
RIGHT = RGBColor(0xFF, 0xE0, 0xC2)  # max right within level (orange)
WIN = RGBColor(0xCB, 0xEA, 0xD4)    # selected max across levels (green)
WIN_B = RGBColor(0x3F, 0x9E, 0x63)
STEP = RGBColor(0xEE, 0xF1, 0xF6)
STEPB = RGBColor(0xC3, 0xCC, 0xD8)
BORDER = RGBColor(0x9A, 0xA5, 0xB1)
GREEN_INK = RGBColor(0x12, 0x34, 0x1F)

OUT_PATH = "/Users/valosek/code/dcm-zurich-lesion-analysis/methodology_spinal_lemniscus_max.pptx"

# ---- worked example (left peaks at C4, right at C5 -> different levels) ----
IN_HEADER = ["slice\n(row)", "vert.\nlevel", "Left tract [%]", "Right tract [%]"]
IN_ROWS = [
    ["8", "C4", "30", "10"],
    ["9", "C4", "55", "12"],
    ["10", "C4", "20", "15"],
    ["14", "C5", "40", "62"],
    ["15", "C5", "48", "44"],
    ["16", "C5", "10", "20"],
]
# per-level maxima in the input table (data-row index, col): 55@r1, 15@r2, 62@r3, 48@r4
IN_HL = {(1, 2): LEFT, (2, 3): RIGHT, (4, 2): LEFT, (3, 3): RIGHT}

LVL_HEADER = ["vert.\nlevel", "Left max\n% (slice)", "Right max\n% (slice)"]
LVL_ROWS = [
    ["C4", "55 (slice 9)", "15 (slice 10)"],
    ["C5", "48 (slice 15)", "62 (slice 14)"],
]
LVL_HL_BASE = {(0, 1): LEFT, (0, 2): RIGHT, (1, 1): LEFT, (1, 2): RIGHT}   # slide 2
LVL_HL_FINAL = {(0, 2): RIGHT, (1, 1): LEFT}                                # slide 3+ (non-winners)
LVL_WIN = {(0, 1): WIN, (1, 2): WIN}                                        # slide 3+ (green winners)


# ---- helpers (all take the target slide's `shapes` collection) ----
def _set_cell_border(cell, color=BORDER, width_pt=0.75):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = tcPr.find(qn(tag))
        if ln is not None:
            tcPr.remove(ln)
        ln = tcPr.makeelement(qn(tag), {"w": str(Emu(Pt(width_pt))), "cap": "flat"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        fill.append(fill.makeelement(qn("a:srgbClr"), {"val": "%02X%02X%02X" % (color[0], color[1], color[2])}))
        ln.append(fill)
        tcPr.append(ln)


def textbox(shapes, x, y, w, h, text, size=12, bold=False, color=INK, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.MIDDLE, italic=False):
    tb = shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(2)
    tf.margin_top = tf.margin_bottom = Pt(1)
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
        r.font.name = "Calibri"
    return tb


def table(shapes, x, y, col_w, row_h, header, rows, highlights=None, win=None, nofill=None):
    highlights = highlights or {}
    win = win or {}
    nofill = nofill or set()
    nrows, ncols = len(rows) + 1, len(header)
    gt = shapes.add_table(nrows, ncols, Inches(x), Inches(y),
                          Inches(sum(col_w)), Inches(row_h * nrows)).table
    gt.first_row = False
    gt.horz_banding = False
    for j, cw in enumerate(col_w):
        gt.columns[j].width = Inches(cw)
    for i in range(nrows):
        gt.rows[i].height = Inches(row_h)

    def fill_cell(cell, txt, fill, bold, border=BORDER, bw=0.75, size=11):
        if fill is None:
            cell.fill.background()   # no fill (transparent)
        else:
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = cell.margin_right = Pt(3)
        cell.margin_top = cell.margin_bottom = Pt(1)
        tf = cell.text_frame
        tf.word_wrap = True
        for k, line in enumerate(str(txt).split("\n")):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = line
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = WHITE if fill == HEADER else INK
            r.font.name = "Calibri"
        _set_cell_border(cell, border, bw)

    for j, htxt in enumerate(header):
        fill_cell(gt.cell(0, j), htxt, HEADER, True)
    for i, rrow in enumerate(rows):
        for j, val in enumerate(rrow):
            if (i, j) in win:
                fill_cell(gt.cell(i + 1, j), val, WIN, True, border=WIN_B, bw=2.0)
            elif (i, j) in nofill:
                fill_cell(gt.cell(i + 1, j), val, None, False)
            else:
                fill = highlights.get((i, j), WHITE)
                fill_cell(gt.cell(i + 1, j), val, fill, (i, j) in highlights)
    return gt


def swatch(shapes, x, y, fill, edge):
    sq = shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.3), Inches(0.3))
    sq.fill.solid()
    sq.fill.fore_color.rgb = fill
    sq.line.color.rgb = edge
    sq.line.width = Pt(1.5)
    sq.shadow.inherit = False


def step_box(shapes, x, y, w, h, num, text):
    box = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = STEP
    box.line.color.rgb = STEPB
    box.line.width = Pt(1.2)
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = "    " + text
    r.font.size = Pt(12.5)
    r.font.color.rgb = INK
    r.font.name = "Calibri"
    # number badge straddling the top-left corner of the box
    d = 0.34
    badge = shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - 0.125), Inches(y - 0.17), Inches(d), Inches(d))
    badge.fill.solid()
    badge.fill.fore_color.rgb = HEADER
    badge.line.fill.background()
    badge.shadow.inherit = False
    bt = badge.text_frame
    bt.vertical_anchor = MSO_ANCHOR.MIDDLE
    bp = bt.paragraphs[0]
    bp.alignment = PP_ALIGN.CENTER
    br = bp.add_run()
    br.text = num
    br.font.size = Pt(13)
    br.font.bold = True
    br.font.color.rgb = WHITE


def arrow(shapes, x0, y0, x1, y1):
    c = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x0), Inches(y0), Inches(x1), Inches(y1))
    c.line.color.rgb = MUTE
    c.line.width = Pt(2.0)
    le = c.line._get_or_add_ln()
    le.append(le.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))


def build_slide(shapes, step):
    """Add all elements that should be visible at build `step` (0..4)."""
    # --- always present ---
    textbox(shapes, 0.5, 0.226, 12.3, 0.398, "Left & right tract lesion damage",
            size=22, bold=True, align=PP_ALIGN.CENTER)
    textbox(shapes, 0.451, 0.937, 4.5, 0.466,
            "Per-slice axial data (one subject) – overlap between the lesion and white matter tracts",
            size=13, bold=True)
    # input per-slice table: highlights appear from step 1; on slide 0 the Left/Right cells are unfilled
    if step >= 1:
        table(shapes, 0.55, 1.5, [0.85, 0.85, 1.25, 1.25], 0.42, IN_HEADER, IN_ROWS, IN_HL)
    else:
        no = {(i, j) for i in range(len(IN_ROWS)) for j in (2, 3)}
        table(shapes, 0.55, 1.5, [0.85, 0.85, 1.25, 1.25], 0.42, IN_HEADER, IN_ROWS, nofill=no)
    # C4 / C5 level bracket labels to the left of the input table
    textbox(shapes, 0.135, 2.326, 0.629, 0.404, "C4", size=18, align=PP_ALIGN.CENTER)
    textbox(shapes, 0.136, 3.596, 0.629, 0.404, "C5", size=18, align=PP_ALIGN.CENTER)
    # legend
    swatch(shapes, 5.4, 1.6, LEFT, BORDER)
    textbox(shapes, 5.8, 1.648, 3.0, 0.205, "max left within level", size=10.5)
    swatch(shapes, 5.4, 2.1, RIGHT, BORDER)
    textbox(shapes, 5.8, 2.148, 3.0, 0.205, "max right within level", size=10.5)
    swatch(shapes, 5.4, 2.6, WIN, WIN_B)
    textbox(shapes, 5.8, 2.58, 3.0, 0.34, "selected max across levels", size=10.5)

    # --- step 1 (from slide 1) ---
    if step >= 1:
        step_box(shapes, 5.35, 3.23, 7.4, 0.7, "1",
                 "Within each vertebral level, take the max across that level's slices. "
                 "Left and right independently.")

    # --- per-level maxima table + first arrow (from slide 2) ---
    if step >= 2:
        if step >= 3:
            table(shapes, 7.0, 4.27, [1.1, 1.5, 1.5], 0.46, LVL_HEADER, LVL_ROWS, LVL_HL_FINAL, LVL_WIN)
        else:
            table(shapes, 7.0, 4.27, [1.1, 1.5, 1.5], 0.46, LVL_HEADER, LVL_ROWS, LVL_HL_BASE)
        arrow(shapes, 9.05, 3.95, 9.05, 4.25)

    # --- step 2 + second arrow (from slide 3) ---
    if step >= 3:
        step_box(shapes, 5.35, 6.05, 7.4, 0.6, "2",
                 "Across vertebral levels, take the max of the per-level maxima. Separately per side.")
        arrow(shapes, 9.05, 5.65, 9.05, 6.05)

    # --- output box + third arrow (from slide 4) ---
    if step >= 4:
        box = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55), Inches(5.511), Inches(4.55), Inches(1.68))
        box.fill.solid()
        box.fill.fore_color.rgb = WIN
        box.line.color.rgb = WIN_B
        box.line.width = Pt(1.6)
        box.shadow.inherit = False
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = tf.margin_right = Pt(8)
        tf.margin_top = Pt(8)
        lines = [
            ("Report per side (one row per subject)", 13, True, GREEN_INK, False),
            ("Left tract 55 %  →  level C4 (slice 9)", 12, False, GREEN_INK, False),
            ("Right tract 62 %  →  level C5 (slice 14)", 12, False, GREEN_INK, False),
            ("Independent lesions → they peak at different levels (C4 vs C5).", 10, False, MUTE, True),
        ]
        for i, (txt, size, bold, color, italic) in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            p.space_before = Pt(0 if i == 0 else (9 if i == 3 else 7))
            r = p.add_run()
            r.text = txt
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.italic = italic
            r.font.color.rgb = color
            r.font.name = "Calibri"
        arrow(shapes, 5.1, 6.35, 5.35, 6.351)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for step in range(5):
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
        build_slide(slide.shapes, step)
    prs.save(OUT_PATH)
    print("saved", OUT_PATH)


if __name__ == "__main__":
    main()
